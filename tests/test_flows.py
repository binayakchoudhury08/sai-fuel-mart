"""End-to-end test against a real Postgres. Run:
DATABASE_URL=postgresql://postgres@/sfm?host=/tmp&port=5433 SETUP_CODE=x python tests/test_flows.py"""
import os
import re
import sys
import uuid
from datetime import timedelta
from decimal import Decimal

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import psycopg2

DSN = os.environ["DATABASE_URL"]
if "supabase" in DSN or "render" in DSN:
    sys.exit("Refusing to run tests against a real database. Use an empty local test database.")
c0 = psycopg2.connect(DSN); c0.autocommit = True
c0.cursor().execute("DROP SCHEMA IF EXISTS pump CASCADE")
c0.close()

from sfm import create_app  # noqa
from sfm.util import today  # noqa

app = create_app()
app.config["TESTING"] = False
T = today()
Y = T - timedelta(days=1)
FAILS = []


def check(cond, msg):
    print(("  ok  " if cond else "  FAIL ") + msg)
    if not cond:
        FAILS.append(msg)


def sql(q, p=None):
    c = psycopg2.connect(DSN)
    cur = c.cursor()
    cur.execute(q, p)
    r = cur.fetchall() if cur.description else None
    c.commit(); c.close()
    return r


class U:
    def __init__(self):
        self.c = app.test_client()

    def token(self):
        with self.c.session_transaction() as s:
            if "_csrf" not in s:
                s["_csrf"] = "t" + uuid.uuid4().hex
            return s["_csrf"]

    def get(self, url, **kw):
        return self.c.get(url, **kw)

    def post(self, url, data=None, ft=None, follow=True):
        d = dict(data or {})
        d["_csrf"] = self.token()
        d["_ft"] = ft or uuid.uuid4().hex
        return self.c.post(url, data=d, follow_redirects=follow)


def flashes(r):
    return re.findall(r'class="flash (\w+)"[^>]*>([^<]+)<', r.get_data(as_text=True))


def has(r, text):
    return text in r.get_data(as_text=True)


owner = U()
print("Setup")
r = owner.get("/login")
check(r.status_code == 302 and "/setup" in r.location, "fresh install redirects to setup")
os.environ["SETUP_CODE"] = "letmein"
r = owner.post("/setup", {"setup_code": "wrong", "username": "boss", "password": "Pump@2026x", "password2": "Pump@2026x"})
check(has(r, "Setup code is wrong"), "wrong setup code refused")
r = owner.post("/setup", {"setup_code": "letmein", "username": "boss", "full_name": "Owner", "password": "Pump@2026x",
                          "password2": "Pump@2026x"})
check(r.status_code == 200 and has(r, "Owner account created"), "owner account created")
check(owner.get("/setup").status_code == 302, "setup closed after first user")

print("Configuration")
owner.post("/settings/rates", {"effective_from": (T - timedelta(days=10)).isoformat(), "rate_MS": "100.00",
                               "rate_HSD": "90.00", "rate_CNG": "80"})
r = owner.post("/settings/tank", {"name": "MS tank", "fuel_code": "MS", "capacity_l": "15000"})
owner.post("/settings/tank", {"name": "HSD tank", "fuel_code": "HSD", "capacity_l": "20000"})
tanks = dict(sql("SELECT fuel_code, id FROM pump.tanks"))
owner.post("/settings/nozzle", {"name": "MS-1", "machine_no": "1", "fuel_code": "MS", "tank_id": tanks["MS"], "sort": 1})
owner.post("/settings/nozzle", {"name": "HSD-1", "machine_no": "1", "fuel_code": "HSD", "tank_id": tanks["HSD"], "sort": 2})
owner.post("/settings/nozzle", {"name": "HSD-2", "machine_no": "2", "fuel_code": "HSD", "tank_id": tanks["HSD"], "sort": 3})
noz = dict(sql("SELECT name, id FROM pump.nozzles"))
check(len(noz) == 3, "3 nozzles configured")
r = owner.post("/settings/nozzle", {"name": "Bad", "fuel_code": "MS", "tank_id": tanks["HSD"]})
check(has(r, "same fuel"), "nozzle on wrong-fuel tank refused")

print("Party + lube")
r = owner.post("/parties/new", {"name": "Sharma Transport", "mobile": "9876543210", "opening_balance": "10000",
                                "opening_date": (T - timedelta(days=5)).isoformat(), "credit_limit": "0"})
pid = sql("SELECT id FROM pump.parties WHERE name='Sharma Transport'")[0][0]
r = owner.post("/parties/new", {"name": "sharma transport "})
check(has(r, "already exists"), "duplicate party name refused")
owner.post("/lube/product", {"name": "Servo 4T 1L", "unit": "pcs", "sale_rate": "350", "opening_stock": "10", "reorder_level": "3"})
lp = sql("SELECT id FROM pump.lube_products")[0][0]

print(f"Day 1 ({Y})")
D = f"/day/{Y.isoformat()}"
r = owner.post(f"{D}/readings", {f"closing_{noz['MS-1']}": "1500"})
check(has(r, "enter the opening reading"), "first-ever reading requires an opening")
r = owner.post(f"{D}/readings", {f"opening_{noz['MS-1']}": "1000", f"closing_{noz['MS-1']}": "1500", f"testing_{noz['MS-1']}": "5",
                                 f"opening_{noz['HSD-1']}": "2000", f"closing_{noz['HSD-1']}": "1999"})
check(has(r, "less than opening"), "closing below opening refused")
check(sql("SELECT COUNT(*) FROM pump.nozzle_readings")[0][0] == 0, "nothing saved when one row is bad (all-or-nothing)")
r = owner.post(f"{D}/readings", {f"opening_{noz['MS-1']}": "1000", f"closing_{noz['MS-1']}": "1500", f"testing_{noz['MS-1']}": "5",
                                 f"opening_{noz['HSD-1']}": "2000", f"closing_{noz['HSD-1']}": "2800",
                                 f"opening_{noz['HSD-2']}": "500", f"closing_{noz['HSD-2']}": "700"})
check(sql("SELECT COUNT(*) FROM pump.nozzle_readings")[0][0] == 3, "3 readings saved")
# MS 495 L x 100 = 49,500 ; HSD 800+200=1000 L x 90 = 90,000 ; fuel = 139,500
modes = dict(sql("SELECT name, id FROM pump.collection_modes"))
owner.post(f"{D}/cash", {f"mode_{modes['PhonePe']}": "20000", f"mode_{modes['Card swipe']}": "5000"})
# credit: 100 L HSD = 9,000 ; driver cash 2,000 ; payment cash 5,000 ; bank payment 3,000 (no drawer effect)
r = owner.post("/parties/txn", {"party": pid, "business_date": Y.isoformat(), "kind": "fuel", "fuel_code": "HSD", "qty": "100",
                                "vehicle_no": "wb 23 a 1234", "challan_no": "C1"})
check(has(r, "Entry SL 1 saved"), "fuel credit via day-page quick form gets SL 1 (no JS)")
owner.post(f"/parties/{pid}/txn", {"business_date": Y.isoformat(), "kind": "cash_advance", "amount": "2000"})
ft = uuid.uuid4().hex
owner.post(f"/parties/{pid}/txn", {"business_date": Y.isoformat(), "kind": "payment", "amount": "5000", "pay_mode": "Cash"}, ft=ft)
r = owner.post(f"/parties/{pid}/txn", {"business_date": Y.isoformat(), "kind": "payment", "amount": "5000", "pay_mode": "Cash"}, ft=ft)
check(has(r, "already saved once"), "double-tap on payment ignored")
check(sql("SELECT COUNT(*) FROM pump.party_txns WHERE kind='payment'")[0][0] == 1, "only ONE payment stored after double submit")
r = owner.post(f"/parties/{pid}/txn", {"business_date": Y.isoformat(), "kind": "payment", "amount": "3000", "pay_mode": "Bank transfer"})
check(has(r, "UTR"), "bank payment needs a reference")
owner.post(f"/parties/{pid}/txn", {"business_date": Y.isoformat(), "kind": "payment", "amount": "3000", "pay_mode": "Bank transfer", "reference": "UTR123"})
rcpts = [x[0] for x in sql("SELECT receipt_no FROM pump.party_txns WHERE kind='payment' ORDER BY id")]
check(len(set(rcpts)) == 2 and all(rcpts), f"payments got unique receipt numbers {rcpts}")
# lube: 2 cash (700), 1 credit (350)
owner.post(f"{D}/lube-sale", {"product_id": lp, "qty": "2", "pay_mode": "Cash"})
owner.post(f"{D}/lube-sale", {"product_id": lp, "qty": "1", "pay_mode": "Credit", "party_id": pid})
r = owner.post(f"{D}/lube-sale", {"product_id": lp, "qty": "50", "pay_mode": "Cash"})
check(has(r, "Only 7"), "cannot sell more lube than in stock")
owner.post(f"{D}/expense", {"description": "Tea", "category": "Tea & snacks", "amount": "150", "from_counter": "1"})
owner.post(f"{D}/expense", {"description": "Electricity bill", "amount": "4000", "from_counter": "0"})
# expected = 139500 + 1050 - (9000+350) - 25000 + 5000 - 2000 - 150 = 109,050
exp = Decimal("109050.00")
r = owner.post(f"{D}/cash", {"den_500": "200", "den_100": "90", "den_50": "1"})  # 100000+9000+50 = 109,050
with app.app_context():
    from sfm import db
    from sfm.services import compute_day, party_balance
    with db.tx() as cur:
        t = compute_day(cur, Y)
        bal = party_balance(cur, pid)
check(t["fuel_total"] == Decimal("139500.00"), f"fuel sale 139,500 (got {t['fuel_total']})")
check(t["expected_cash"] == exp, f"expected cash 1,09,050 (got {t['expected_cash']})")
check(t["counted_cash"] == exp and t["difference"] == 0, f"note count = 1,09,050, difference 0 (got {t['difference']})")
# balance = 10000 + 9000 + 2000 + 350 - 5000 - 3000 = 13,350
check(bal == Decimal("13350.00"), f"party balance 13,350 (got {bal})")
r = owner.get(D)
check(r.status_code == 200 and has(r, "₹1,09,050.00"), "day page renders with Indian formatting")

print("Workflow + locks")
r = owner.post(f"{D}/submit", {})
check(has(r, "does not match Credit transport") or has(r, "Credit does not match"), "submit blocked until closing credit matches Credit transport")
r = owner.post(f"{D}/credit", {f"credit_{pid}": "10000"})
check(has(r, "do not match"), "wrong closing total shows a mismatch")
r = owner.post(f"{D}/credit", {f"credit_{pid}": "11000"})
check(has(r, "Every transporter matches"), "closing total 11,000 = slips 9,000 + cash 2,000 → matched")
r = owner.post(f"{D}/submit", {})
check(sql("SELECT status FROM pump.days WHERE business_date=%s", (Y,))[0][0] == "submitted", "day submitted")
owner.post(f"{D}/lock", {})
check(sql("SELECT status FROM pump.days WHERE business_date=%s", (Y,))[0][0] == "locked", "day locked")
r = owner.post(f"/parties/{pid}/txn", {"business_date": Y.isoformat(), "kind": "payment", "amount": "1", "pay_mode": "Cash"})
check(has(r, "locked"), "owner cannot add to locked day through app")
try:
    sql("INSERT INTO pump.expenses (business_date, description, amount) VALUES (%s,'sneaky',1)", (Y,))
    check(False, "database blocks direct insert on locked day")
except psycopg2.Error as e:
    check("locked" in str(e), "database itself blocks direct insert on locked day")
try:
    sql("DELETE FROM pump.party_txns WHERE id=(SELECT MIN(id) FROM pump.party_txns)")
    check(False, "database blocks deleting party entries")
except psycopg2.Error:
    check(True, "database blocks deleting party entries")
try:
    sql("UPDATE pump.audit_log SET summary='x'")
    check(False, "audit log is read-only")
except psycopg2.Error:
    check(True, "audit log cannot be edited")
r = owner.post(f"{D}/unlock", {"reason": ""})
check(has(r, "Write why"), "unlock needs a reason")
owner.post(f"{D}/unlock", {"reason": "Correcting test"})
check(sql("SELECT status FROM pump.days WHERE business_date=%s", (Y,))[0][0] == "open", "owner unlocked with reason")
check(sql("SELECT COUNT(*) FROM pump.audit_log WHERE action='unlock' AND summary LIKE %s", ("%Correcting test%",))[0][0] == 1,
      "unlock reason recorded in audit log")

print(f"Day 2 ({T}) – openings carried forward")
D2 = f"/day/{T.isoformat()}"
r = owner.get(D2)
check(has(r, 'value="1500.00" readonly'), "today's MS-1 opening = yesterday's closing, read-only")
owner.post(f"{D2}/readings", {f"closing_{noz['MS-1']}": "1600", f"closing_{noz['HSD-1']}": "2900", f"closing_{noz['HSD-2']}": "750"})
op = sql("SELECT opening FROM pump.nozzle_readings WHERE business_date=%s AND nozzle_id=%s", (T, noz["MS-1"]))[0][0]
check(op == Decimal("1500.00"), "server used carried-forward opening")

print("Operator restrictions")
owner.post("/users/new", {"username": "ravi", "full_name": "Ravi", "role": "operator", "password": "Ravi@12345"})
op_ = U()
op_.post("/login", {"username": "ravi", "password": "Ravi@12345"})
r = op_.get("/dashboard")
check(r.status_code == 302 and "password" in r.location, "new user forced to change password")
op_.post("/account/password", {"old_password": "Ravi@12345", "new_password": "Ravi#2026a", "new_password2": "Ravi#2026a"})
check(op_.get("/dashboard").status_code == 200, "operator signed in after password change")
r = op_.post(f"{D2}/readings", {f"opening_{noz['MS-1']}": "1400", f"closing_{noz['MS-1']}": "1600"})
check(has(r, "Only the owner can override"), "operator cannot change opening reading")
tid = sql("SELECT id FROM pump.party_txns WHERE kind='payment' ORDER BY id LIMIT 1")[0][0]
r = op_.post(f"/party-txn/{tid}/void", {"reason": "x"}, follow=False)
check(r.status_code == 403, "operator cannot cancel entries")
r = op_.post(f"/parties/{pid}/txn", {"business_date": (T - timedelta(days=5)).isoformat(), "kind": "payment", "amount": "10", "pay_mode": "Cash"})
check(has(r, "Only the owner can change entries older"), "operator cannot back-date beyond window")
r = op_.post(f"/parties/{pid}/txn", {"business_date": (T + timedelta(days=1)).isoformat(), "kind": "payment", "amount": "10", "pay_mode": "Cash"})
check(has(r, "future"), "nobody can enter future dates")
check(op_.get("/settings").status_code == 403 and op_.get("/audit").status_code == 403, "operator blocked from settings & audit")
r = op_.post(f"/parties/{pid}/txn", {"kind": "payment", "amount": "abc", "pay_mode": "Cash"})
check(has(r, "must be a number"), "junk amount rejected, not turned into 0")
bad = U()
for i in range(5):
    bad.post("/login", {"username": "ravi", "password": "wrong"})
r = bad.post("/login", {"username": "ravi", "password": "Ravi#2026a"})
check(has(r, "Try again in"), "5 wrong passwords lock the account for 15 minutes")
r = op_.c.post(f"/parties/{pid}/txn", data={"kind": "payment", "amount": "10", "pay_mode": "Cash", "_csrf": "forged"})
check(r.status_code == 400, "forged form (bad CSRF) rejected")

print("Void + statement")
owner.post(f"/party-txn/{tid}/void", {"reason": "Wrong party"})
with app.app_context():
    with db.tx() as cur:
        bal2 = party_balance(cur, pid)
check(bal2 == Decimal("18350.00"), f"cancelling ₹5,000 payment restores balance to 18,350 (got {bal2})")
r = owner.get(f"/parties/{pid}?from={(T - timedelta(days=10)).isoformat()}")
check(has(r, "Wrong party") and has(r, 'class="void"'), "cancelled payment stays visible, struck through")
lm = sql("SELECT id FROM pump.lube_moves WHERE pay_mode='Credit'")[0][0]
owner.post(f"/lube-move/{lm}/void", {"reason": "returned"})
with app.app_context():
    with db.tx() as cur:
        bal3 = party_balance(cur, pid)
check(bal3 == Decimal("18000.00"), f"cancelling credit lube sale also removes it from party (got {bal3})")
stock = sql("SELECT SUM(qty) FROM pump.lube_moves WHERE NOT voided")[0][0]
check(stock == Decimal("8.00"), f"lube stock back to 8 (got {stock})")

print("Exports + pages")
for url in ["/dashboard", D, D2, f"/day/{Y.isoformat()}/print", "/parties", f"/parties/{pid}", "/lube", "/tanks",
            "/staff", "/staff/salary", "/reports", "/reports/outstanding", "/reports/credit-register", "/settings",
            "/audit", "/users", "/proofs", "/admin/import-legacy", f"/party-txn/{rcpts[1] and sql('SELECT id FROM pump.party_txns WHERE receipt_no=%s',(rcpts[1],))[0][0]}/receipt"]:
    r = owner.get(url)
    check(r.status_code == 200, f"GET {url} → {r.status_code}")
for url, mime in [(f"/parties/{pid}/bill.pdf?from={(T - timedelta(days=10)).isoformat()}", "application/pdf"),
                  (f"/parties/{pid}/bill.xlsx", "spreadsheet"), ("/reports?format=xlsx", "spreadsheet"),
                  ("/reports/outstanding?format=xlsx", "spreadsheet"), ("/reports/credit-register?format=xlsx", "spreadsheet"),
                  ("/reports/lube.xlsx", "spreadsheet"), ("/reports/backup.xlsx", "spreadsheet")]:
    r = owner.get(url)
    check(r.status_code == 200 and mime in r.mimetype and len(r.data) > 1000, f"download {url.split('?')[0]} ({len(r.data)} bytes)")
open("/tmp/bill.pdf", "wb").write(owner.get(f"/parties/{pid}/bill.pdf?from={(T - timedelta(days=10)).isoformat()}").data)

print("Tanks + ATG + staff + rate change")
owner.post("/tanks/dip", {"business_date": Y.isoformat(), f"volume_{tanks['MS']}": "8000", f"volume_{tanks['HSD']}": "12000"})
owner.post("/tanks/receipt", {"business_date": T.isoformat(), "tank_id": tanks["HSD"], "invoice_no": "INV9", "invoice_qty_l": "4000",
                              "received_qty_l": "3980"})
r = owner.post("/tanks/receipt", {"business_date": T.isoformat(), "tank_id": tanks["HSD"], "invoice_no": "INV9", "received_qty_l": "1"})
check(has(r, "already recorded"), "same tanker invoice cannot be entered twice")
owner.post("/tanks/dip", {"business_date": T.isoformat(), f"volume_{tanks['HSD']}": "15800"})
with app.app_context():
    from sfm.services import tank_reconciliation
    with db.tx() as cur:
        rec = {r["tank"]["fuel_code"]: r for r in tank_reconciliation(cur, T)}
# HSD book = 12000 + 3980 - (100 + 50) = 15830 ; dip 15800 → -30
check(rec["HSD"]["book"] == Decimal("15830.00") and rec["HSD"]["variance"] == Decimal("-30.00"),
      f"tank variance HSD −30 L (got book {rec['HSD']['book']}, var {rec['HSD']['variance']})")
os.environ["TANK_SYNC_SECRET"] = "s3"
r = app.test_client().post("/api/tank-sync", json={"secret": "nope", "tanks": []})
check(r.status_code == 403, "ATG sync rejects wrong secret")
r = app.test_client().post("/api/tank-sync", json={"secret": "s3", "tanks": [{"fuel_type": "MS", "net_volume": "7500.5", "dip_mm": "1200"}]})
check(r.status_code == 200 and r.json["synced"] == ["MS"], f"ATG sync accepted (old bookmarklet format) {r.json}")
owner.post("/staff/save", {"emp_code": "EMP001", "name": "Raju", "monthly_salary": "12000"})
sid = sql("SELECT id FROM pump.staff")[0][0]
owner.post("/staff/attendance", {"work_date": T.isoformat(), f"st_{sid}": "P"})
owner.post("/staff/payment", {"staff_id": sid, "kind": "advance", "amount": "500", "pay_mode": "Cash from counter",
                              "business_date": T.isoformat(), "for_month": T.strftime("%Y-%m")})
with app.app_context():
    with db.tx() as cur:
        t2 = compute_day(cur, T)
check(t2["staff_cash"] == Decimal("500.00"), "staff advance from counter reduces today's expected cash")
owner.post("/settings/rates", {"effective_from": T.isoformat(), "rate_MS": "101"})
r1 = sql("SELECT rate FROM pump.nozzle_readings WHERE business_date=%s AND nozzle_id=%s", (T, noz["MS-1"]))[0][0]
r0 = sql("SELECT rate FROM pump.nozzle_readings WHERE business_date=%s AND nozzle_id=%s", (Y, noz["MS-1"]))[0][0]
check(r1 == Decimal("101.00") and r0 == Decimal("100.00"), "rate change re-prices today only; yesterday keeps old rate")

print("Credit transport register (restored v1 features)")
r = owner.post("/parties/new", {"name": "SVT Logistics", "mobile": "9830011111", "vehicles": "WB23A1111", "opening_balance": "0"})
svt = sql("SELECT id FROM pump.parties WHERE name='SVT Logistics'")[0][0]
r = owner.post("/transport/entry", {"business_date": T.isoformat(), "party": svt, "fuel_code": "HSD", "qty": "100",
                                    "cash_taken": "500", "challan_no": "C9", "vehicle_no": "WB23A1111", "slip_no": "S1"})
first = sql("SELECT sl_no, debit, cash_taken, qty FROM pump.party_txns WHERE party_id=%s AND NOT voided ORDER BY id", (svt,))
check(first and first[0][1] == Decimal("9500.00") and first[0][2] == Decimal("500.00"),
      f"one register row = 100 L x 90 + cash 500 = 9,500 (got {first})")
r = owner.post("/transport/entry", {"business_date": T.isoformat(), "party": svt, "fuel_code": "HSD", "qty": "10", "slip_no": "S1"})
check(sql("SELECT COUNT(*) FROM pump.party_txns WHERE slip_no='S1' AND NOT voided")[0][0] == 2, "same slip allowed for many entries")
sid_extra = sql("SELECT id FROM pump.party_txns WHERE slip_no='S1' AND qty=10")[0][0]
owner.post(f"/party-txn/{sid_extra}/void", {"reason": "test entry"})
owner.post("/transport/entry", {"business_date": T.isoformat(), "party": svt, "fuel_code": "HSD", "qty": "50", "slip_no": "S2"})
owner.post("/transport/entry", {"business_date": T.isoformat(), "party": svt, "cash_taken": "300", "slip_no": "S3"})
sls = [x[0] for x in sql("SELECT sl_no FROM pump.party_txns WHERE business_date=%s AND kind IN ('fuel','cash_advance') ORDER BY sl_no", (T,))]
check(sls == sorted(set(sls)) and len(sls) >= 3, f"per-day SL numbers are unique and ordered {sls}")
with app.app_context():
    with db.tx() as cur:
        t3 = compute_day(cur, T)
check(t3["cash_advance"] == 0, "no separate 'cash given to drivers' line in daily closing")
check(t3["credit_fuel"] == Decimal("14300.00"), f"credit given = full slip totals 9,500 + 4,500 + 300 (got {t3['credit_fuel']})")
tid2 = sql("SELECT id FROM pump.party_txns WHERE slip_no='S2' AND NOT voided")[0][0]
sl2 = sql("SELECT sl_no FROM pump.party_txns WHERE id=%s", (tid2,))[0][0]
owner.post(f"/transport/entry/{tid2}/correct", {"qty": "60", "fuel_code": "HSD", "slip_no": "S2", "reason": "meter slip said 60"})
newr = sql("SELECT sl_no, qty, corrected_from FROM pump.party_txns WHERE slip_no='S2' AND NOT voided")
check(newr[0][0] == sl2 and newr[0][1] == Decimal("60.00") and newr[0][2] == tid2, "correction keeps the same SL No and links to the old entry")
check(sql("SELECT voided FROM pump.party_txns WHERE id=%s", (tid2,))[0][0], "corrected entry stays visible as cancelled")
import openpyxl, io, zipfile, subprocess
q = f"party={svt}&period=day&d={T.isoformat()}"
r = owner.get(f"/transport/download?{q}&fmt=xlsx")
wbx = openpyxl.load_workbook(io.BytesIO(r.data)); ws = wbx.active
vals = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
head = [c.value for c in ws[6]]
check(head == ["Date", "SL", "Vehicle No.", "Slip No.", "Challan No.", "HSD Qty (L)", "Rate", "HSD Amount", "Cash Taken", "Total"],
      f"day Excel bill has the simple columns {head}")
tot = [row for row in ws.iter_rows(values_only=True) if row[0] == "TOTAL"][0]
check(tot[5] == 160 and tot[8] == 800 and tot[9] == 15200, f"TOTAL row: 160 L, cash 800, total 15,200 {tot}")
check(not any(isinstance(v, str) and ("Balance" in v or "Net Payable" in v or "Discount" in v) for v in vals),
      "Excel bill shows no balance, discount or net payable")
r = owner.get(f"/transport/download?{q}&fmt=pdf")
open("/tmp/transport_bill.pdf", "wb").write(r.data)
txt = subprocess.run(["pdftotext", "-layout", "/tmp/transport_bill.pdf", "-"], capture_output=True, text=True).stdout
check(r.data[:4] == b"%PDF" and "Vehicle No." in txt and "TOTAL" in txt and "15,200.00" in txt, "PDF day bill with columns and total")
check("Balance" not in txt and "Net" not in txt and "Discount" not in txt and "Due" not in txt, "PDF bill shows no balance due / net payable / discount")
r = owner.get(f"/transport/download?party={svt}&period=month&month={T.strftime('%Y-%m')}&fmt=zip")
names = zipfile.ZipFile(io.BytesIO(r.data)).namelist()
mon = T.strftime("%b %Y")
check(f"SVT Logistics/SVT Logistics Bill {mon}.pdf" in names and f"SVT Logistics/SVT Logistics Bill {mon}.xlsx" in names
      and any(n.startswith("SVT Logistics/Daily PDF/") for n in names) and any(n.startswith("SVT Logistics/Daily Excel/") for n in names),
      f"month ZIP: month bill PDF+Excel and daily PDF+Excel {names}")
r = owner.get(f"/transport/download?period=month&month={T.strftime('%Y-%m')}&fmt=zip")
folders = sorted({n.split("/")[0] for n in zipfile.ZipFile(io.BytesIO(r.data)).namelist()})
check("SVT Logistics" in folders and "Sharma Transport" in folders, f"all-transporters ZIP: separate folder each {folders}")
r = owner.get(f"/transport/download?period=month&month={T.strftime('%Y-%m')}&fmt=pdf")
zn = zipfile.ZipFile(io.BytesIO(r.data)).namelist() if r.mimetype == "application/zip" else []
check(len(zn) == 2 and all(n.endswith(".pdf") for n in zn), f"all transporters + PDF → one PDF per transporter {zn}")
r = owner.get(f"/transport/download?party={svt}&period=month&month={T.strftime('%Y-%m')}&fmt=pdf")
check(r.data[:4] == b"%PDF" and mon in r.headers.get("Content-Disposition", ""), "month PDF bill named by month")
r = owner.get(f"/transport/bill.pdf?party={svt}&from={T.isoformat()}&to={T.isoformat()}", follow_redirects=True)
check(r.data[:4] == b"%PDF", "old bill links still work")
r = owner.get(f"/transport/download?party={svt}&period=day&d={(T - timedelta(days=9)).isoformat()}&fmt=pdf", follow_redirects=True)
check(has(r, "No transport entries"), "empty day says so instead of an empty file")
r = owner.post("/parties/txn", {"party": svt, "kind": "payment", "amount": "1000", "pay_mode": "Cash", "pay_for": "Lube"})
check(sql("SELECT pay_for FROM pump.party_txns WHERE party_id=%s AND kind='payment'", (svt,))[0][0] == "Lube", "payment remembers Fuel/Lube")

print("Restored reports")
for url in ["/transport", f"/transport?d={T.isoformat()}&q=WB23", "/reports/nozzles", "/reports/digital",
            "/reports/tanks", "/analytics", "/proofs", f"/proofs?d={T.isoformat()}"]:
    check(owner.get(url).status_code == 200, f"GET {url}")
for url in ["/reports?format=pdf", "/reports/sales.pdf", "/reports/nozzles?format=pdf", "/reports/nozzles?format=xlsx",
            "/reports/digital?format=pdf", "/reports/digital?format=xlsx", "/reports/tanks?format=pdf", "/reports/tanks?format=xlsx",
            "/reports/outstanding?format=pdf", "/reports/credit-register?format=pdf",
            f"/transport/register.pdf?party={svt}&from={T.isoformat()}&to={T.isoformat()}",
            f"/transport/register.xlsx?party={svt}&from={T.isoformat()}&to={T.isoformat()}", "/reports/attendance.xlsx", "/reports/business.xlsx",
            f"/proofs/report.pdf?d={T.isoformat()}"]:
    r = owner.get(url)
    check(r.status_code == 200 and len(r.data) > 900, f"download {url} ({len(r.data)} bytes)")
wbb = openpyxl.load_workbook(io.BytesIO(owner.get("/reports/business.xlsx").data))
check(len(wbb["Dashboard"]._charts) >= 1 and len(wbb["Transport summary"]._charts) == 1, "business Excel has line + pie charts")
wba = openpyxl.load_workbook(io.BytesIO(owner.get("/reports/attendance.xlsx").data))
check("Summary" in wba.sheetnames and any("Raju" in n for n in wba.sheetnames), f"attendance Excel has summary + per-staff sheets {wba.sheetnames[:4]}")

print("Tanker receipt with compartments")
r = owner.post("/tanks/receipt", {"business_date": T.isoformat(), "invoice_no": "INV77", "order_no": "OR1", "carrier_name": "IOC Tptr",
                                  "challan_density": "835", "observed_density": "830.5", "invoice_HSD": "4000", "invoice_MS": "8000",
                                  "c1_fuel": "MS", "c1_cdip": "1200", "c1_odip": "1195", "c1_vol": "4000",
                                  "c2_fuel": "MS", "c2_cdip": "1200", "c2_odip": "1202", "c2_vol": "3990",
                                  "c3_fuel": "HSD", "c3_cdip": "1100", "c3_odip": "1100", "c3_vol": "4000"})
rr = sql("SELECT t.fuel_code, r.received_qty_l, r.dip_variance_mm, r.receipt_group FROM pump.fuel_receipts r JOIN pump.tanks t ON t.id=r.tank_id WHERE invoice_no='INV77' ORDER BY 1")
check(len(rr) == 2 and rr[1][1] == Decimal("7990.00") and rr[1][2] == Decimal("-3.00") and rr[0][3] == rr[1][3],
      f"one tanker → MS 7,990 L (dip var −3 mm) + HSD 4,000 L, linked {rr}")
check(sql("SELECT COUNT(*) FROM pump.audit_log WHERE summary LIKE %s", ("%INV77%DENSITY ALERT%",))[0][0] >= 1, "density difference > 3 flagged")
r = owner.get(f"/tanks?d={T.isoformat()}")
check(has(r, "Density alerts") and has(r, "Tankers this month"), "tank page shows month tanker stats")

print()
print("FAILED:" if FAILS else "ALL CHECKS PASSED", len(FAILS) or "")
for f in FAILS:
    print(" -", f)
sys.exit(1 if FAILS else 0)
