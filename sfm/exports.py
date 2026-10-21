"""PDF / Excel builders. Pure functions: data in, BytesIO out."""
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .services import PARTY_KIND_LABEL
from .util import inr, now_ist

INK = colors.HexColor("#1B2A41")
RULE = colors.HexColor("#C9CED6")
TINT = colors.HexColor("#EEF1F5")


def _particulars(r):
    k = r["kind"]
    if k == "fuel":
        return f"{r['fuel_code']} {r['qty']} {r.get('unit') or 'L'} @ {r['rate']}"
    if k == "lube":
        return f"{r.get('lube_name') or r.get('remarks') or 'Lube'} × {r['qty']}"
    if k == "payment":
        s = f"Received – {r['pay_mode']}"
        if r.get("reference"):
            s += f" ({r['reference']})"
        if r.get("receipt_no"):
            s += f" #{r['receipt_no']}"
        return s
    s = PARTY_KIND_LABEL.get(k, k)
    if r.get("remarks") and k not in ("fuel",):
        s += f" – {r['remarks']}"
    return s


def party_bill_pdf(p, stmt, start, end, st):
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=14 * mm,
                            title=f"Statement – {p['name']}", author=st.get("station_name", ""))
    ss = getSampleStyleSheet()
    h = ParagraphStyle("h", parent=ss["Title"], fontName="Helvetica-Bold", fontSize=17, textColor=INK,
                       alignment=0, spaceAfter=1 * mm)
    small = ParagraphStyle("s", parent=ss["Normal"], fontSize=8.5, textColor=colors.HexColor("#4A5568"),
                           leading=11)
    body = ParagraphStyle("b", parent=ss["Normal"], fontSize=10, leading=13)
    cell = ParagraphStyle("c", parent=ss["Normal"], fontSize=8.5, leading=10.5)
    story = [Paragraph(st.get("station_name") or "Statement", h)]
    sub = " | ".join(x for x in [st.get("station_address"), st.get("station_phone") and f"Ph {st.get('station_phone')}",
                                 st.get("station_gstin") and f"GSTIN {st.get('station_gstin')}"] if x)
    if sub:
        story.append(Paragraph(sub, small))
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph(f"<b>Account statement:</b> {p['name']}"
                           + (f" &nbsp; Mobile {p['mobile']}" if p.get("mobile") else "")
                           + (f" &nbsp; GSTIN {p['gstin']}" if p.get("gstin") else ""), body))
    story.append(Paragraph(f"Period {start.strftime('%d %b %Y')} to {end.strftime('%d %b %Y')}", small))
    story.append(Spacer(1, 3 * mm))

    data = [["Date", "Particulars", "Vehicle", "Challan / Slip", "Debit (Rs)",
             "Credit (Rs)", "Balance (Rs)"]]
    data.append([start.strftime("%d-%m-%y"), "Opening balance", "", "", "", "", inr(stmt["opening"], symbol=False)])
    for r in stmt["rows"]:
        if r["voided"]:
            continue
        cs = " / ".join(x for x in [r.get("challan_no"), r.get("slip_no")] if x)
        data.append([r["business_date"].strftime("%d-%m-%y"), Paragraph(_particulars(r), cell),
                     r.get("vehicle_no") or "", cs,
                     inr(r["debit"], symbol=False) if r["debit"] else "",
                     inr(r["credit"], symbol=False) if r["credit"] else "",
                     inr(r["running"], symbol=False)])
    tot = stmt["totals"]
    data.append(["", "Total for period", "", "", inr(tot["debit"], symbol=False),
                 inr(tot["credit"], symbol=False), ""])
    data.append(["", "Closing balance (amount due)" if stmt["closing"] >= 0 else "Closing balance (advance)",
                 "", "", "", "", inr(stmt["closing"], symbol=False)])
    widths = [22 * mm, 92 * mm, 30 * mm, 38 * mm, 30 * mm, 30 * mm, 31 * mm]
    t = Table(data, colWidths=widths, repeatRows=1)
    style = [
        ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.5),
        ("FONT", (0, 1), (-1, -1), "Helvetica", 8.5),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, 0), (-1, 0), INK),
        ("ALIGN", (4, 0), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LINEBELOW", (0, 1), (-1, -3), 0.3, RULE),
        ("BACKGROUND", (0, -2), (-1, -1), TINT),
        ("FONT", (0, -2), (-1, -1), "Helvetica-Bold", 9),
        ("FONT", (0, 1), (-1, 1), "Helvetica-Oblique", 8.5),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    t.setStyle(TableStyle(style))
    story.append(t)
    story.append(Spacer(1, 5 * mm))
    story.append(Paragraph(f"Fuel supplied in period: {tot['fuel_qty']} L. "
                           "Please check and report any difference within 7 days.", small))

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#6B7280"))
        canvas.drawString(12 * mm, 7 * mm, f"Generated {now_ist().strftime('%d %b %Y %I:%M %p')}")
        canvas.drawRightString(landscape(A4)[0] - 12 * mm, 7 * mm, f"Page {d.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    buf.seek(0)
    return buf


HEAD_FILL = PatternFill("solid", start_color="1B2A41")
HEAD_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="C9CED6")


def _excel_value(v):
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, datetime):
        return v.replace(tzinfo=None)
    return v


def write_sheet(ws, headers, rows, money_cols=(), widths=None):
    ws.append(headers)
    for c in ws[ws.max_row]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append([_excel_value(v) for v in r])
    for col in money_cols:
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            for c in row:
                c.number_format = "#,##,##0.00"
    for i, h in enumerate(headers, 1):
        w = (widths or {}).get(i) or max(10, min(45, max((len(str(r[i - 1])) for r in rows), default=0) + 2,
                                               len(str(h)) + 2))
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def party_bill_xlsx(p, stmt, start, end, st):
    wb = Workbook()
    ws = wb.active
    ws.title = "Statement"
    ws.append([st.get("station_name", "")])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Statement: {p['name']}  |  {start.strftime('%d %b %Y')} to {end.strftime('%d %b %Y')}"])
    ws.append([])
    rows = [[start, "Opening balance", "", "", "", None, None, stmt["opening"]]]
    for r in stmt["rows"]:
        if r["voided"]:
            continue
        rows.append([r["business_date"], _particulars(r), r.get("vehicle_no") or "", r.get("challan_no") or "",
                     r.get("slip_no") or "", r["debit"] or None, r["credit"] or None, r["running"]])
    rows.append([None, "Total", "", "", "", stmt["totals"]["debit"], stmt["totals"]["credit"], None])
    rows.append([None, "Closing balance", "", "", "", None, None, stmt["closing"]])
    start_row = ws.max_row + 1
    write_sheet(ws, ["Date", "Particulars", "Vehicle", "Challan", "Slip", "Debit", "Credit", "Balance"],
                rows, money_cols=(6, 7, 8), widths={1: 12, 2: 42})
    for row in ws.iter_rows(min_row=start_row + 1, min_col=1, max_col=1):
        for c in row:
            if isinstance(c.value, (date, datetime)):
                c.number_format = "dd-mm-yyyy"
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def workbook_bytes(sheets):
    """sheets: list of (title, headers, rows, money_cols)"""
    wb = Workbook()
    wb.remove(wb.active)
    for title, headers, rows, money_cols in sheets:
        ws = wb.create_sheet(title[:31])
        write_sheet(ws, headers, rows, money_cols)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, (date, datetime)) and not isinstance(c.value, datetime):
                    c.number_format = "dd-mm-yyyy"
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ====================================================================== credit transport bill (v1 format)
import re as _re
import zipfile as _zip
from decimal import ROUND_HALF_UP as _RHU

GREEN = colors.HexColor("#16A34A")


def bill_filename(party_name, start, ext):
    """'SVT Logistics Bill 110626.pdf' — same naming as the old app."""
    safe = _re.sub(r'[\\/:*?"<>|]', "", party_name or "Transporter").strip() or "Transporter"
    return f"{safe} Bill {start.strftime('%d%m%y')}.{ext}"


def bill_number(party, start, end):
    initials = "".join(w[0] for w in _re.findall(r"[A-Za-z0-9]+", party["name"]))[:4].upper() or "TR"
    return f"{initials}-{start.strftime('%d%m%y')}" + (f"-{end.strftime('%d%m%y')}" if end != start else "")


def bill_figures(party, rows):
    live = [r for r in rows if not r["voided"]]
    litres = sum((Decimal(r["qty"]) for r in live if r["kind"] == "fuel"), Decimal(0))
    fuel_amt = sum((Decimal(r["fuel_amount"]) for r in live), Decimal(0))
    cash = sum((r["cash"] for r in live), Decimal(0))
    total = sum((Decimal(r["debit"]) for r in live), Decimal(0))
    rate = Decimal(party.get("discount_per_litre") or 0)
    discount = (rate * litres).quantize(Decimal("0.01"), _RHU)
    return {"rows": live, "litres": litres, "fuel_amount": fuel_amt, "cash": cash, "total": total,
            "disc_rate": rate, "discount": discount, "net": total - discount}


def _letterhead(st, title, meta_lines, width):
    ss = getSampleStyleSheet()
    name = ParagraphStyle("n", parent=ss["Title"], fontName="Helvetica-Bold", fontSize=19, leading=22,
                          textColor=INK, alignment=0, spaceAfter=0)
    sub = ParagraphStyle("s", parent=ss["Normal"], fontSize=8.5, leading=12, textColor=colors.HexColor("#475467"))
    bt = ParagraphStyle("bt", parent=ss["Normal"], fontName="Helvetica-Bold", fontSize=12.5, leading=16,
                        textColor=colors.white, alignment=2)
    bm = ParagraphStyle("bm", parent=ss["Normal"], fontSize=8.5, leading=12, textColor=colors.white, alignment=2)
    left = [Paragraph((st.get("station_name") or "SAI FUEL MART").upper(), name)]
    if st.get("station_address"):
        left.append(Paragraph(st["station_address"], sub))
    bits = [x for x in [st.get("station_gstin") and f"GSTIN: {st['station_gstin']}",
                        st.get("station_phone") and f"Ph: {st['station_phone']}"] if x]
    if bits:
        left.append(Paragraph(" | ".join(bits), sub))
    right = [Paragraph(title, bt)] + [Paragraph(m, bm) for m in meta_lines]
    t = Table([[left, right]], colWidths=[width * 0.64, width * 0.36])
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("BACKGROUND", (1, 0), (1, 0), INK),
                           ("LEFTPADDING", (0, 0), (0, 0), 0), ("LEFTPADDING", (1, 0), (1, 0), 14),
                           ("RIGHTPADDING", (1, 0), (1, 0), 14), ("TOPPADDING", (0, 0), (-1, -1), 9),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 9)]))
    rule = Table([[""]], colWidths=[width], rowHeights=[3.5])
    rule.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), GREEN)]))
    return [t, Spacer(1, 3), rule, Spacer(1, 10)]


def _footer_fn(st, label, pagesize):
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#E5E7EB"))
        canvas.line(20, 30, pagesize[0] - 20, 30)
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#667085"))
        canvas.drawString(20, 17, f"{st.get('station_name') or 'Sai Fuel Mart'} — {label} — computer generated "
                                  f"{now_ist().strftime('%d-%m-%Y %I:%M %p')}")
        canvas.drawRightString(pagesize[0] - 20, 17, f"Page {doc.page}")
        canvas.restoreState()
    return footer


def transport_bill_pdf(party, rows, start, end, st):
    f = bill_figures(party, rows)
    buf = BytesIO()
    ps = landscape(A4)
    width = ps[0] - 40
    doc = SimpleDocTemplate(buf, pagesize=ps, leftMargin=20, rightMargin=20, topMargin=20, bottomMargin=38,
                            title=f"Bill {party['name']}")
    ss = getSampleStyleSheet()
    lab = ParagraphStyle("l", parent=ss["Normal"], fontSize=9.5, leading=13, textColor=INK)
    period = start.strftime("%d-%m-%Y") + ("" if start == end else " to " + end.strftime("%d-%m-%Y"))
    el = _letterhead(st, "CREDIT TRANSPORT BILL",
                     [f"Bill No: {bill_number(party, start, end)}", f"Period: {period}",
                      f"Generated: {now_ist().strftime('%d-%m-%Y %H:%M')}"], width)
    el.append(Paragraph(f"<b>Bill To:</b> {party['name']}", lab))
    if party.get("vehicles"):
        el.append(Paragraph(f"<b>Vehicle No:</b> {party['vehicles']}", lab))
    if party.get("mobile"):
        el.append(Paragraph(f"<b>Contact:</b> {party['mobile']}", lab))
    if party.get("gstin"):
        el.append(Paragraph(f"<b>GSTIN:</b> {party['gstin']}", lab))
    el.append(Spacer(1, 10))
    data = [["Date", "SL", "Challan", "Vehicle", "Slip", "Fuel", "Qty (L)", "Rate", "Fuel Amt", "Cash Taken", "Final Amt"]]
    for r in f["rows"]:
        data.append([r["business_date"].strftime("%d-%m-%Y"), str(r["sl_no"] or ""), r["challan_no"] or "-",
                     r["vehicle_no"] or "-", r["slip_no"] or "-", r["fuel_code"] or "Cash",
                     f"{Decimal(r['qty']):.2f}" if r["kind"] == "fuel" else "-",
                     f"{Decimal(r['rate']):.2f}" if r["kind"] == "fuel" else "-",
                     inr(r["fuel_amount"], symbol=False), inr(r["cash"], symbol=False),
                     inr(r["debit"], symbol=False)])
    if not f["rows"]:
        data.append(["No entries in this period"] + [""] * 10)
    data.append(["", "", "", "", "", "TOTAL", f"{f['litres']:.2f}", "-", inr(f["fuel_amount"], symbol=False),
                 inr(f["cash"], symbol=False), inr(f["total"], symbol=False)])
    ti = len(data) - 1
    style = [("BACKGROUND", (0, 0), (-1, 0), INK), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
             ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8),
             ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
             ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
             ("ALIGN", (1, 1), (1, -1), "CENTER"), ("ALIGN", (6, 1), (-1, -1), "RIGHT"),
             ("ROWBACKGROUNDS", (0, 1), (-1, ti - 1), [colors.white, colors.HexColor("#F8FAFC")]),
             ("BACKGROUND", (0, ti), (-1, ti), colors.HexColor("#DCFCE7")), ("FONTNAME", (0, ti), (-1, ti), "Helvetica-Bold")]
    if f["disc_rate"] > 0:
        data.append(["Total amount"] + [""] * 9 + [inr(f["total"], symbol=False)])
        data.append([f"Less discount (Rs {f['disc_rate']}/L x {f['litres']:.2f} L)"] + [""] * 9 + [inr(f["discount"], symbol=False)])
        data.append(["NET PAYABLE"] + [""] * 9 + [inr(f["net"], symbol=False)])
        a, b, c = ti + 1, ti + 2, ti + 3
        for i in (a, b, c):
            style += [("SPAN", (0, i), (9, i)), ("ALIGN", (0, i), (9, i), "RIGHT"), ("FONTNAME", (0, i), (-1, i), "Helvetica-Bold")]
        style += [("BACKGROUND", (0, a), (-1, a), colors.HexColor("#F1F5F9")),
                  ("BACKGROUND", (0, b), (-1, b), colors.HexColor("#FEE2E2")), ("TEXTCOLOR", (0, b), (-1, b), colors.HexColor("#B91C1C")),
                  ("BACKGROUND", (0, c), (-1, c), GREEN), ("TEXTCOLOR", (0, c), (-1, c), colors.white),
                  ("FONTSIZE", (0, c), (-1, c), 10.5), ("TOPPADDING", (0, c), (-1, c), 8), ("BOTTOMPADDING", (0, c), (-1, c), 8)]
    widths = [62, 28, 74, 80, 62, 40, 62, 52, 88, 80, 90]
    scale = width / sum(widths)
    t = Table(data, colWidths=[w * scale for w in widths], repeatRows=1)
    t.setStyle(TableStyle(style))
    el += [t, Spacer(1, 34)]
    sig = Table([["", ""], ["Prepared By", "Authorised Signatory"]], colWidths=[width / 2, width / 2])
    sig.setStyle(TableStyle([("LINEABOVE", (0, 1), (0, 1), 0.6, colors.HexColor("#94A3B8")),
                             ("LINEABOVE", (1, 1), (1, 1), 0.6, colors.HexColor("#94A3B8")),
                             ("FONTSIZE", (0, 1), (-1, 1), 8.5), ("TEXTCOLOR", (0, 1), (-1, 1), colors.HexColor("#667085")),
                             ("ALIGN", (1, 1), (1, 1), "RIGHT")]))
    el.append(sig)
    fn = _footer_fn(st, "Credit Transport Bill", ps)
    doc.build(el, onFirstPage=fn, onLaterPages=fn)
    buf.seek(0)
    return buf


def transport_bill_xlsx(party, rows, start, end, st):
    f = bill_figures(party, rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "Transport Bill"
    ws.append([st.get("station_name") or "Sai Fuel Mart"]); ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Credit Transport Bill {bill_number(party, start, end)} — {party['name']} — "
               f"{start.strftime('%d-%m-%Y')} to {end.strftime('%d-%m-%Y')}"])
    ws.append([])
    head = ["Date", "SL No", "Party Name", "Challan No", "Vehicle No", "Slip No", "Fuel", "Qty", "Rate",
            "Fuel Amount", "Cash Taken", "Final Amount"]
    ws.append(head)
    hr = ws.max_row
    for c in ws[hr]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    for r in f["rows"]:
        ws.append([r["business_date"], r["sl_no"], party["name"], r["challan_no"], r["vehicle_no"], r["slip_no"],
                   r["fuel_code"] or "Cash", float(r["qty"]) if r["kind"] == "fuel" else None,
                   float(r["rate"]) if r["kind"] == "fuel" else None, float(r["fuel_amount"]), float(r["cash"]),
                   float(r["debit"])])
        ws.cell(row=ws.max_row, column=1).number_format = "dd-mm-yyyy"
    ws.append([])
    ws.append(["", "", "", "", "", "", "TOTAL", float(f["litres"]), None, float(f["fuel_amount"]), float(f["cash"]),
               float(f["total"])])
    tot = ws.max_row
    for c in ws[tot]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", start_color="D9EAD3")
    if f["disc_rate"] > 0:
        for label, val, fill, font in [
                ("Total Amount", f["total"], "F1F5F9", Font(bold=True)),
                (f"Discount (Rs {f['disc_rate']}/L x {f['litres']:.2f} L)", f["discount"], "FEE2E2", Font(bold=True, color="B91C1C")),
                ("Net Payable", f["net"], "16A34A", Font(bold=True, size=13, color="FFFFFF"))]:
            ws.append([""] * 10 + [label, float(val)])
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", start_color=fill); c.font = font
    for row in ws.iter_rows(min_row=hr + 1, min_col=8, max_col=12):
        for c in row:
            c.number_format = "#,##,##0.00"
    for i, w in enumerate([12, 7, 24, 12, 13, 10, 7, 10, 9, 14, 13, 15], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.column_dimensions["K"].width = 34
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    buf = BytesIO(); wb.save(buf); buf.seek(0)
    return buf


def daily_bills_zip(parties_rows, st):
    """parties_rows: list of (party, rows). One PDF + one Excel per party per day,
    in folders 'Month YYYY/PDF' and 'Month YYYY/Excel' (per party when several)."""
    buf = BytesIO()
    used = set()
    multi = len(parties_rows) > 1
    with _zip.ZipFile(buf, "w", _zip.ZIP_DEFLATED) as zf:
        for party, rows in parties_rows:
            by_day = {}
            for r in rows:
                if not r["voided"]:
                    by_day.setdefault(r["business_date"], []).append(r)
            for d, day_rows in sorted(by_day.items()):
                base = (f"{_re.sub(r'[^A-Za-z0-9 ]+', '', party['name']).strip()}/" if multi else "") + d.strftime("%B %Y")
                for sub, data, ext in (("PDF", transport_bill_pdf(party, day_rows, d, d, st), "pdf"),
                                       ("Excel", transport_bill_xlsx(party, day_rows, d, d, st), "xlsx")):
                    name = f"{base}/{sub}/{bill_filename(party['name'], d, ext)}"
                    n = 2
                    while name in used:
                        name = name.rsplit(".", 1)[0] + f"_{n}.{ext}"; n += 1
                    used.add(name)
                    zf.writestr(name, data.getvalue())
    buf.seek(0)
    return buf, len(used) // 2


# ====================================================================== generic report PDF
def report_pdf(st, title, subtitle, headers, rows, right_cols=(), totals=None, widths=None, landscape_mode=True,
               notes=None):
    buf = BytesIO()
    ps = landscape(A4) if landscape_mode else A4
    width = ps[0] - 40
    doc = SimpleDocTemplate(buf, pagesize=ps, leftMargin=20, rightMargin=20, topMargin=20, bottomMargin=38,
                            title=title)
    ss = getSampleStyleSheet()
    cell = ParagraphStyle("c", parent=ss["Normal"], fontSize=7.6, leading=9.4)
    el = _letterhead(st, title.upper(), [subtitle, f"Generated: {now_ist().strftime('%d-%m-%Y %H:%M')}"], width)

    def fmt(v):
        if v is None:
            return ""
        if isinstance(v, Decimal):
            return inr(v, symbol=False)
        if isinstance(v, (date,)) and not isinstance(v, datetime):
            return v.strftime("%d-%m-%Y")
        if isinstance(v, str) and len(v) > 28:
            return Paragraph(v, cell)
        return str(v)
    data = [headers] + [[fmt(v) for v in r] for r in rows]
    if not rows:
        data.append(["No data for this selection"] + [""] * (len(headers) - 1))
    if totals:
        data.append([fmt(v) for v in totals])
    if widths:
        sc = width / sum(widths)
        cw = [w * sc for w in widths]
    else:
        cw = [width / len(headers)] * len(headers)
    t = Table(data, colWidths=cw, repeatRows=1)
    style = [("BACKGROUND", (0, 0), (-1, 0), INK), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
             ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 7.6),
             ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D1D5DB")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
             ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
             ("ROWBACKGROUNDS", (0, 1), (-1, -1 if not totals else -2), [colors.white, colors.HexColor("#F8FAFC")])]
    for c in right_cols:
        style.append(("ALIGN", (c, 1), (c, -1), "RIGHT"))
    if totals:
        style += [("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#DCFCE7")), ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold")]
    t.setStyle(TableStyle(style))
    el.append(t)
    for n in notes or []:
        el += [Spacer(1, 6), Paragraph(n, ParagraphStyle("nt", parent=ss["Normal"], fontSize=8.5, leading=11))]
    fn = _footer_fn(st, title, ps)
    doc.build(el, onFirstPage=fn, onLaterPages=fn)
    buf.seek(0)
    return buf
