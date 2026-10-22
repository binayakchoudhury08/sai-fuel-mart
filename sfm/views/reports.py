from datetime import timedelta
from decimal import Decimal

from flask import Blueprint, render_template, request, send_file

from .. import db
from ..exports import workbook_bytes
from ..security import audit, login_required
from ..services import PARTY_KIND_LABEL, fuels, lube_stock, party_balances, range_summary, rates_for
from ..util import ZERO, UserError, now_ist, parse_date, today


def _deliver_compat(buf, mimetype=None, as_attachment=True, download_name=None, **kw):
    from ..storage import deliver
    return deliver(buf, download_name, mimetype)


bp = Blueprint("reports", __name__)
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _range(default_days=None):  # noqa: C901
    t = today()
    s = request.args.get("from")
    e = request.args.get("to")
    start = parse_date(s, "From date") if s else (t - timedelta(days=default_days) if default_days else t.replace(day=1))
    end = parse_date(e, "To date") if e else t
    if end < start:
        start, end = end, start
    if (end - start).days > 400:
        raise UserError("Choose a range of 400 days or less.")
    return start, end


@bp.route("/reports")
@login_required("reports")
def index():
    start, end = _range()
    with db.tx() as cur:
        rows = range_summary(cur, start, end)
        fl = fuels(cur, active_only=False)
    keys = ["fuel_total", "lube_total", "gross_sale", "digital", "credit", "payments", "cash_advance",
            "expenses", "staff_cash", "expected_cash"]
    tot = {k: sum((r[k] for r in rows), ZERO) for k in keys}
    tot["qty"] = {f["code"]: sum((r["qty"].get(f["code"], ZERO) for r in rows), ZERO) for f in fl}
    tot["difference"] = sum((r["difference"] for r in rows if r["difference"] is not None), ZERO)
    tot["short_days"] = sum(1 for r in rows if r["difference"] is not None and r["difference"] < 0)
    if request.args.get("format") == "xlsx":
        heads = ["Date", "Status"] + [f"{f['code']} qty" for f in fl] + [
            "Fuel sale", "Lube sale", "Total sale", "Digital", "Credit given", "Payments recd",
            "Expenses", "Staff cash", "Expected cash", "Counted cash", "Short(-)/Excess(+)"]
        data = [[r["date"], r["status"]] + [r["qty"].get(f["code"], ZERO) for f in fl] +
                [r["fuel_total"], r["lube_total"], r["gross_sale"], r["digital"], r["credit"], r["payments"],
                 r["expenses"], r["staff_cash"], r["expected_cash"], r["counted"],
                 r["difference"]] for r in rows]
        n = len(fl)
        money_cols = tuple(range(3, 3 + n + 11))
        return _deliver_compat(workbook_bytes([("Daily sales", heads, data, money_cols)]), mimetype=XLSX,
                         as_attachment=True, download_name=f"Sales {start:%d%m%y}-{end:%d%m%y}.xlsx")
    if request.args.get("format") == "pdf":
        return sales_pdf()
    return render_template("reports/index.html", rows=rows, tot=tot, start=start, end=end, fuels=fl)


@bp.route("/reports/outstanding")
@login_required("reports")
def outstanding():
    with db.tx() as cur:
        rows = party_balances(cur)
    rows = sorted(rows, key=lambda r: -Decimal(r["balance"]))
    t = today()
    for r in rows:
        r["days_since_payment"] = (t - r["last_payment"]).days if r["last_payment"] else None
    if request.args.get("format") == "xlsx":
        data = [[r["name"], r["mobile"], "Active" if r["active"] else "Inactive", r["total_debit"],
                 r["total_credit"], r["balance"], r["credit_limit"] or None, r["last_payment"],
                 r["days_since_payment"]] for r in rows]
        return _deliver_compat(workbook_bytes([("Outstanding", ["Party", "Mobile", "Status", "Total debit",
                                          "Total credit", "Balance", "Credit limit", "Last payment",
                                          "Days since payment"], data, (4, 5, 6, 7))]),
                         mimetype=XLSX, as_attachment=True, download_name=f"Outstanding {t:%d%m%y}.xlsx")
    if request.args.get("format") == "pdf":
        return outstanding_pdf()
    return render_template("reports/outstanding.html", rows=rows)


@bp.route("/reports/credit-register")
@login_required("reports")
def credit_register():
    start, end = _range()
    kind = request.args.get("kind") or ""
    fmt = request.args.get("format")
    try:
        page = max(1, int(request.args.get("page") or 1))
    except ValueError:
        page = 1
    per = 300
    with db.tx() as cur:
        where = "t.business_date BETWEEN %s AND %s"
        params = [start, end]
        if kind in PARTY_KIND_LABEL:
            where += " AND t.kind=%s"
            params.append(kind)
        cur.execute(f"""SELECT COUNT(*) AS n, COALESCE(SUM(debit) FILTER (WHERE NOT voided),0) AS d,
                               COALESCE(SUM(credit) FILTER (WHERE NOT voided),0) AS c
                        FROM party_txns t WHERE {where}""", params)
        agg = cur.fetchone()
        sql = f"""SELECT t.*, p.name AS party_name, lp.name AS lube_name FROM party_txns t
                  JOIN parties p ON p.id=t.party_id LEFT JOIN lube_products lp ON lp.id=t.lube_product_id
                  WHERE {where} ORDER BY t.business_date, t.id"""
        if fmt != "xlsx":
            sql += f" LIMIT {per} OFFSET {(page - 1) * per}"
        cur.execute(sql, params)
        rows = cur.fetchall()
    tot = {"debit": Decimal(agg["d"]), "credit": Decimal(agg["c"])}
    pages = max(1, -(-agg["n"] // per))
    if fmt == "xlsx":
        data = [[r["business_date"], r["party_name"], PARTY_KIND_LABEL[r["kind"]], r["fuel_code"] or r["lube_name"],
                 r["qty"] or None, r["rate"] or None, r["debit"] or None, r["credit"] or None, r["pay_mode"],
                 r["receipt_no"], r["vehicle_no"], r["challan_no"], r["slip_no"], r["reference"], r["remarks"],
                 "CANCELLED: " + (r["void_reason"] or "") if r["voided"] else "", r["created_by"]] for r in rows]
        heads = ["Date", "Party", "Type", "Item", "Qty", "Rate", "Debit", "Credit", "Mode", "Receipt", "Vehicle",
                 "Challan", "Slip", "Reference", "Remarks", "Cancelled", "Entered by"]
        return _deliver_compat(workbook_bytes([("Credit register", heads, data, (6, 7, 8))]), mimetype=XLSX,
                         as_attachment=True, download_name=f"Credit register {start:%d%m%y}-{end:%d%m%y}.xlsx")
    if request.args.get("format") == "pdf":
        return credit_register_pdf()
    return render_template("reports/credit_register.html", rows=rows, tot=tot, start=start, end=end, kind=kind,
                           KIND=PARTY_KIND_LABEL, page=page, pages=pages, count=agg["n"])


@bp.route("/reports/lube.xlsx")
@login_required("reports")
def lube_xlsx():
    with db.tx() as cur:
        prods = lube_stock(cur)
        cur.execute("""SELECT m.*, p.name FROM lube_moves m JOIN lube_products p ON p.id=m.product_id
                       ORDER BY m.business_date, m.id""")
        moves = cur.fetchall()
    s1 = [[p["name"], p["unit"], p["sale_rate"], p["stock"], Decimal(p["stock"]) * Decimal(p["sale_rate"]),
           p["reorder_level"], "Active" if p["active"] else "Inactive"] for p in prods]
    s2 = [[m["business_date"], m["name"], m["kind"], m["qty"], m["rate"], m["amount"], m["pay_mode"], m["remarks"],
           "CANCELLED" if m["voided"] else "", m["created_by"]] for m in moves]
    return _deliver_compat(workbook_bytes([
        ("Stock", ["Product", "Unit", "Sale rate", "Stock", "Stock value", "Reorder at", "Status"], s1, (3, 5)),
        ("Movements", ["Date", "Product", "Type", "Qty", "Rate", "Amount", "Mode", "Remarks", "Cancelled", "By"],
         s2, (5, 6))]), mimetype=XLSX, as_attachment=True, download_name=f"Lube stock {today():%d%m%y}.xlsx")


BACKUP_TABLES = ["settings", "users", "fuels", "fuel_rates", "tanks", "nozzles", "days", "nozzle_readings",
                 "collection_modes", "day_collections", "expenses", "parties", "party_txns", "lube_products",
                 "lube_moves", "tank_dips", "fuel_receipts", "staff", "attendance", "staff_payments", "proofs",
                 "audit_log"]


@bp.route("/reports/backup.xlsx")
@login_required("backup")
def backup():
    sheets = []
    with db.tx() as cur:
        for t in BACKUP_TABLES:
            cur.execute(f"SELECT * FROM {t} ORDER BY 1")
            rows = cur.fetchall()
            cols = [c.name for c in cur.description]
            if t == "users":
                cols = [c for c in cols if c != "password_hash"]
            data = [[(str(r[c]) if isinstance(r[c], (dict, list)) else r[c]) for c in cols] for r in rows]
            sheets.append((t, cols, data, ()))
        audit(cur, "export", "backup", None, "Downloaded full backup")
    return _deliver_compat(workbook_bytes(sheets), mimetype=XLSX, as_attachment=True,
                     download_name=f"SaiFuelMart backup {now_ist():%Y-%m-%d %H%M}.xlsx")


# ====================================================================== v2.1: restored reports
from datetime import date as _date
from flask import abort, redirect, url_for
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Font
from io import BytesIO
from ..exports import report_pdf, write_sheet
from ..services import settings_map, tank_reconciliation, lube_stock as _lube_stock
from ..util import month_bounds, money, ZERO as _Z
import calendar


def _pdf(st, title, start, end, heads, rows, right, totals=None, widths=None, fname=None, notes=None):
    return _deliver_compat(report_pdf(st, title, f"{start:%d-%m-%Y} to {end:%d-%m-%Y}" if start else "", heads, rows,
                                right_cols=right, totals=totals, widths=widths, notes=notes),
                     mimetype="application/pdf", as_attachment=True,
                     download_name=(fname or title) + (f" {start:%d%m%y}-{end:%d%m%y}" if start else "") + ".pdf")


@bp.route("/reports/sales.pdf")
@login_required("reports")
def sales_pdf():
    start, end = _range()
    with db.tx() as cur:
        rows = range_summary(cur, start, end)
        st = settings_map(cur)
    heads = ["Date", "Status", "MS L", "HSD L", "CNG kg", "Total sale", "Digital", "Credit", "Received",
             "Expenses", "Expected cash", "Counted", "Short/Excess"]
    data = [[r["date"], r["status"], r["qty"].get("MS"), r["qty"].get("HSD"), r["qty"].get("CNG"), r["gross_sale"],
             r["digital"], r["credit"], r["payments"], r["expenses"], r["expected_cash"], r["counted"],
             r["difference"]] for r in rows]
    s = lambda k: sum((r[k] for r in rows), _Z)
    q = lambda c: sum((r["qty"].get(c, _Z) for r in rows), _Z)
    tot = ["TOTAL", "", q("MS"), q("HSD"), q("CNG"), s("gross_sale"), s("digital"), s("credit"), s("payments"),
           s("expenses"), s("expected_cash"), "", sum((r["difference"] for r in rows if r["difference"] is not None), _Z)]
    return _pdf(st, "Daily sales and cash", start, end, heads, data, range(2, 13), tot,
                [55, 45, 50, 55, 45, 70, 62, 62, 62, 58, 70, 62, 62], "Sales report")


@bp.route("/reports/nozzles")
@login_required("reports")
def nozzles():
    start, end = _range()
    fuel = request.args.get("fuel") or ""
    with db.tx() as cur:
        sql = """SELECT r.*, n.name, n.machine_no, n.fuel_code, (r.closing-r.opening-r.testing) AS sale,
                        ROUND((r.closing-r.opening-r.testing)*r.rate, 2) AS amount
                 FROM nozzle_readings r JOIN nozzles n ON n.id=r.nozzle_id
                 WHERE r.business_date BETWEEN %s AND %s"""
        p = [start, end]
        if fuel:
            sql += " AND n.fuel_code=%s"; p.append(fuel)
        cur.execute(sql + " ORDER BY r.business_date DESC, n.sort, n.name", p)
        rows = cur.fetchall()
        cur.execute(sql.replace("SELECT r.*, n.name, n.machine_no, n.fuel_code, (r.closing-r.opening-r.testing) AS sale,\n                        ROUND((r.closing-r.opening-r.testing)*r.rate, 2) AS amount",
                                "SELECT n.name, n.fuel_code, SUM(r.closing-r.opening-r.testing) AS sale, SUM(r.testing) AS testing, SUM(ROUND((r.closing-r.opening-r.testing)*r.rate,2)) AS amount, COUNT(*) AS days")
                    + " GROUP BY n.name, n.fuel_code, n.sort ORDER BY n.sort, n.name", p)
        per_nozzle = cur.fetchall()
        fl = fuels(cur, active_only=False)
        st = settings_map(cur)
    tot_sale = sum((Decimal(r["sale"]) for r in rows), _Z)
    tot_amt = sum((Decimal(r["amount"]) for r in rows), _Z)
    fmt = request.args.get("format")
    heads = ["Date", "Nozzle", "DU", "Fuel", "Opening", "Closing", "Testing", "Sale", "Rate", "Amount"]
    data = [[r["business_date"], r["name"], r["machine_no"], r["fuel_code"], r["opening"], r["closing"], r["testing"],
             r["sale"], r["rate"], r["amount"]] for r in rows]
    if fmt == "xlsx":
        return _deliver_compat(workbook_bytes([
            ("Nozzle readings", heads, data, (5, 6, 7, 8, 9, 10)),
            ("Per nozzle", ["Nozzle", "Fuel", "Days", "Sale", "Testing", "Amount"],
             [[n["name"], n["fuel_code"], n["days"], n["sale"], n["testing"], n["amount"]] for n in per_nozzle], (4, 5, 6))]),
            mimetype=XLSX, as_attachment=True, download_name=f"Nozzle report {start:%d%m%y}-{end:%d%m%y}.xlsx")
    if fmt == "pdf":
        return _pdf(st, "Nozzle readings", start, end, heads, data, range(4, 10),
                    ["TOTAL", "", "", "", "", "", "", tot_sale, "", tot_amt], [55, 70, 35, 35, 70, 70, 50, 60, 45, 70],
                    "Nozzle report")
    return render_template("reports/nozzles.html", rows=rows, per_nozzle=per_nozzle, start=start, end=end,
                           fuel=fuel, fuels=fl, tot_sale=tot_sale, tot_amt=tot_amt)


@bp.route("/reports/digital")
@login_required("reports")
def digital():
    start, end = _range()
    with db.tx() as cur:
        cur.execute("SELECT * FROM collection_modes ORDER BY sort, name")
        modes = cur.fetchall()
        cur.execute("""SELECT business_date, mode_id, amount FROM day_collections
                       WHERE business_date BETWEEN %s AND %s AND amount <> 0""", (start, end))
        grid = {}
        for r in cur.fetchall():
            grid.setdefault(r["business_date"], {})[r["mode_id"]] = Decimal(r["amount"])
        st = settings_map(cur)
    used = [m for m in modes if any(m["id"] in g for g in grid.values()) or m["active"]]
    dates = sorted(grid, reverse=True)
    rows = [[d] + [grid[d].get(m["id"], _Z) for m in used] + [sum(grid[d].values(), _Z)] for d in dates]
    totals = ["TOTAL"] + [sum((grid[d].get(m["id"], _Z) for d in dates), _Z) for m in used] + \
             [sum((sum(g.values(), _Z) for g in grid.values()), _Z)]
    heads = ["Date"] + [m["name"] for m in used] + ["Total"]
    fmt = request.args.get("format")
    if fmt == "xlsx":
        return _deliver_compat(workbook_bytes([("Digital collections", heads, rows + [totals], tuple(range(2, len(heads) + 1)))]),
                         mimetype=XLSX, as_attachment=True, download_name=f"Digital collections {start:%d%m%y}-{end:%d%m%y}.xlsx")
    if fmt == "pdf":
        return _pdf(st, "Digital collections", start, end, heads, rows, range(1, len(heads)), totals, None,
                    "Digital collections")
    return render_template("reports/digital.html", heads=heads, rows=rows, totals=totals, start=start, end=end)


def tank_range(cur, start, end):
    """Day-by-day book stock vs dip per tank, newest first, with ₹ value of gain/loss."""
    from ..services import tank_series
    rows = [r for r in tank_series(cur, start, end)
            if r["actual"] is not None or r["received"] or r["sold"]]
    return sorted(rows, key=lambda r: (r["date"], r["tank"]["name"]), reverse=True)


@bp.route("/reports/tanks")
@login_required("reports")
def tanks():
    start, end = _range(default_days=13) if not request.args.get("from") else _range()
    if (end - start).days > 186:
        raise UserError("Tank report: choose 6 months or less.")
    with db.tx() as cur:
        rows = tank_range(cur, start, end)
        cur.execute("""SELECT r.*, t.name AS tank_name FROM fuel_receipts r JOIN tanks t ON t.id=r.tank_id
                       WHERE r.business_date BETWEEN %s AND %s AND NOT r.voided ORDER BY r.business_date""", (start, end))
        receipts = cur.fetchall()
        st = settings_map(cur)
    heads = ["Date", "Tank", "Opening", "Tanker", "Meter sale", "Book stock", "Dip stock", "Gain/Loss L", "% of sale", "Gain/Loss ₹"]
    data = [[r["date"], r["tank"]["name"], r["opening"], r["received"], r["sold"], r["book"], r["actual"], r["variance"],
             r["variance_pct"], r["value"]] for r in rows]
    tl = sum((r["variance"] for r in rows if r["variance"] is not None), _Z)
    tv = sum((r["value"] for r in rows if r["value"] is not None), _Z)
    fmt = request.args.get("format")
    if fmt == "xlsx":
        rheads = ["Date", "Tank", "Invoice", "Vehicle", "Order no", "Carrier", "Invoice L", "Received L", "Short L",
                  "Challan density", "Observed density", "Dip variance mm", "Water", "Remarks"]
        rdata = [[r["business_date"], r["tank_name"], r["invoice_no"], r["vehicle_no"], r["order_no"], r["carrier_name"],
                  r["invoice_qty_l"], r["received_qty_l"],
                  (Decimal(r["invoice_qty_l"]) - Decimal(r["received_qty_l"])) if r["invoice_qty_l"] else None,
                  r["challan_density"], r["observed_density"], r["dip_variance_mm"], r["water_checked"], r["remarks"]]
                 for r in receipts]
        return _deliver_compat(workbook_bytes([("Stock variance", heads, data, (3, 4, 5, 6, 7, 8, 10)),
                                         ("Tanker receipts", rheads, rdata, (7, 8, 9))]),
                         mimetype=XLSX, as_attachment=True, download_name=f"Tank report {start:%d%m%y}-{end:%d%m%y}.xlsx")
    if fmt == "pdf":
        return _pdf(st, "Tank stock and variance", start, end, heads, data, range(2, 10),
                    ["TOTAL", "", "", "", "", "", "", tl, "", tv], None, "Tank report")
    return render_template("reports/tanks.html", rows=rows, receipts=receipts, start=start, end=end, tl=tl, tv=tv)


@bp.route("/reports/outstanding.pdf")
@login_required("reports")
def outstanding_pdf():
    with db.tx() as cur:
        rows = sorted(party_balances(cur), key=lambda r: -Decimal(r["balance"]))
        st = settings_map(cur)
    t = today()
    data = [[r["name"], r["mobile"], r["opening_bal"], r["fuel_credit"], r["lube_credit"], r["received"], r["discounts"],
             r["balance"], r["last_payment"], (t - r["last_payment"]).days if r["last_payment"] else ""] for r in rows]
    s = lambda k: sum((Decimal(r[k]) for r in rows), _Z)
    return _deliver_compat(report_pdf(st, "Outstanding balances", f"As on {t:%d-%m-%Y}",
                                ["Party", "Mobile", "Opening", "Fuel credit", "Lube credit", "Received", "Discount",
                                 "Balance", "Last paid", "Days"], data, right_cols=range(2, 10),
                                totals=["TOTAL", "", s("opening_bal"), s("fuel_credit"), s("lube_credit"), s("received"),
                                        s("discounts"), s("balance"), "", ""],
                                widths=[130, 70, 65, 70, 65, 70, 60, 75, 60, 35]),
                     mimetype="application/pdf", as_attachment=True, download_name=f"Outstanding {t:%d%m%y}.pdf")


@bp.route("/reports/credit-register.pdf")
@login_required("reports")
def credit_register_pdf():
    start, end = _range()
    with db.tx() as cur:
        cur.execute("""SELECT t.*, p.name AS party_name FROM party_txns t JOIN parties p ON p.id=t.party_id
                       WHERE t.business_date BETWEEN %s AND %s AND NOT t.voided ORDER BY t.business_date, t.id""",
                    (start, end))
        rows = cur.fetchall()
        st = settings_map(cur)
    data = [[r["business_date"], r["sl_no"] or "", r["party_name"], PARTY_KIND_LABEL[r["kind"]],
             (f"{r['qty']} {r['fuel_code']}" if r["kind"] == "fuel" else (r["pay_mode"] or "")) + (f" #{r['receipt_no']}" if r["receipt_no"] else ""),
             r["vehicle_no"], r["debit"] or None, r["credit"] or None] for r in rows]
    return _pdf(st, "Party credit register", start, end,
                ["Date", "SL", "Party", "Type", "Details", "Vehicle", "Debit", "Credit"], data, (6, 7),
                ["TOTAL", "", "", "", "", "", sum((Decimal(r["debit"]) for r in rows), _Z),
                 sum((Decimal(r["credit"]) for r in rows), _Z)], [55, 25, 120, 80, 130, 70, 70, 70], "Credit register")


# ------------------------------------------------------------------ attendance workbook (old export restored)
@bp.route("/reports/attendance.xlsx")
@login_required("reports")
def attendance_xlsx():
    ym = request.args.get("m") or today().strftime("%Y-%m")
    from .staff import salary_sheet
    first, last = month_bounds(ym)
    with db.tx() as cur:
        sheet = salary_sheet(cur, ym)
        cur.execute("""SELECT a.*, s.name FROM attendance a JOIN staff s ON s.id=a.staff_id
                       WHERE a.work_date BETWEEN %s AND %s ORDER BY a.work_date""", (first, last))
        marks = cur.fetchall()
        cur.execute("""SELECT sp.*, s.name FROM staff_payments sp JOIN staff s ON s.id=sp.staff_id
                       WHERE sp.for_month=%s ORDER BY sp.business_date""", (ym,))
        pays = cur.fetchall()
    days = [(first + timedelta(days=i)) for i in range((last - first).days + 1)]
    grid = {}
    for m in marks:
        grid.setdefault(m["staff_id"], {})[m["work_date"]] = m["status"]
    summary = [[r["staff"]["emp_code"], r["staff"]["name"], r["staff"]["designation"], r["staff"]["monthly_salary"],
                r["att"].get("P", 0), r["att"].get("H", 0), r["att"].get("O", 0), r["att"].get("L", 0),
                r["att"].get("A", 0), sheet["ndays"] - r["marked"], r["paid_days"], r["earned"], r["advance"],
                r["salary_paid"], r["bonus"], r["balance"]] for r in sheet["rows"]]
    details = [[r["staff"]["name"]] + [grid.get(r["staff"]["id"], {}).get(d, "") for d in days] for r in sheet["rows"]]
    sheets = [("Summary", ["Code", "Name", "Role", "Monthly salary", "Present", "Half", "Off", "Leave", "Absent",
                           "Unmarked", "Paid days", "Earned", "Advance", "Salary paid", "Bonus", "Still to pay"],
               summary, (4, 12, 13, 14, 15, 16)),
              ("Attendance grid", ["Staff"] + [str(d.day) for d in days], details, ()),
              ("Payments", ["Date", "Staff", "Type", "Mode", "Amount", "Cancelled", "Remarks"],
               [[p["business_date"], p["name"], p["kind"], p["pay_mode"], p["amount"], "YES" if p["voided"] else "",
                 p["remarks"]] for p in pays], (5,))]
    for r in sheet["rows"]:
        sid = r["staff"]["id"]
        sheets.append((f"{r['staff']['name'][:24]} {r['staff']['emp_code'][-4:]}", ["Date", "Day", "Status"],
                       [[d, d.strftime("%a"), {"P": "Present", "H": "Half day", "A": "Absent", "L": "Leave",
                                              "O": "Weekly off"}.get(grid.get(sid, {}).get(d, ""), "")] for d in days], ()))
    return _deliver_compat(workbook_bytes(sheets), mimetype=XLSX, as_attachment=True,
                     download_name=f"Attendance {calendar.month_name[first.month]} {first.year}.xlsx")


# ------------------------------------------------------------------ analytics
@bp.route("/analytics")
@login_required("reports")
def analytics():
    ym = request.args.get("m") or today().strftime("%Y-%m")
    try:
        first, last = month_bounds(ym)
    except Exception:
        abort(404)
    last = min(last, today())
    with db.tx() as cur:
        rows = range_summary(cur, first, last)
        parties = party_balances(cur)
        prods = _lube_stock(cur)
        cur.execute("""SELECT p.name, SUM(-m.qty) AS q, SUM(m.amount) AS a FROM lube_moves m
                       JOIN lube_products p ON p.id=m.product_id
                       WHERE m.kind='sale' AND NOT m.voided AND m.business_date BETWEEN %s AND %s
                       GROUP BY p.name ORDER BY q DESC LIMIT 5""", (first, last))
        top_lube = cur.fetchall()
        cur.execute("""SELECT p.name, SUM(t.debit) AS credit, SUM(t.qty) FILTER (WHERE t.kind='fuel') AS litres
                       FROM party_txns t JOIN parties p ON p.id=t.party_id
                       WHERE NOT t.voided AND t.kind IN ('fuel','cash_advance','lube') AND t.business_date BETWEEN %s AND %s
                       GROUP BY p.name ORDER BY credit DESC LIMIT 6""", (first, last))
        top_parties = cur.fetchall()
        cur.execute("SELECT COUNT(*) AS n FROM staff WHERE active")
        staff_n = cur.fetchone()["n"]
        cur.execute("SELECT status, COUNT(*) AS n FROM attendance WHERE work_date=%s GROUP BY status", (today(),))
        att_today = {r["status"]: r["n"] for r in cur.fetchall()}
        cur.execute("SELECT COALESCE(SUM(amount),0) AS a FROM staff_payments WHERE for_month=%s AND NOT voided", (ym,))
        salary_paid = Decimal(cur.fetchone()["a"])
        cur.execute("""SELECT category, SUM(amount) AS a FROM expenses WHERE NOT voided AND business_date BETWEEN %s AND %s
                       GROUP BY category ORDER BY a DESC""", (first, last))
        exp_cat = cur.fetchall()
        fl = fuels(cur, active_only=False)
    s = lambda k: sum((r[k] for r in rows), _Z)
    k = {"sale": s("gross_sale"), "fuel": s("fuel_total"), "lube": s("lube_total"), "credit": s("credit"),
         "received": s("payments"), "expenses": s("expenses"), "digital": s("digital"),
         "diff": sum((r["difference"] for r in rows if r["difference"] is not None), _Z),
         "days": len([r for r in rows if r["gross_sale"]]),
         "qty": {f["code"]: sum((r["qty"].get(f["code"], _Z) for r in rows), _Z) for f in fl}}
    k["avg"] = (k["sale"] / k["days"]).quantize(Decimal("1")) if k["days"] else _Z
    stock_value = sum((Decimal(p["stock"]) * Decimal(p["sale_rate"]) for p in prods if p["active"]), _Z)
    low = sum(1 for p in prods if p["active"] and Decimal(p["stock"]) <= Decimal(p["reorder_level"]))
    outstanding = sum((Decimal(p["balance"]) for p in parties if Decimal(p["balance"]) > 0), _Z)
    peak = max((r["gross_sale"] for r in rows), default=_Z) or Decimal(1)
    y, m = (int(x) for x in ym.split("-"))
    prev_m = f"{y - (m == 1)}-{12 if m == 1 else m - 1:02d}"
    next_m = f"{y + (m == 12)}-{1 if m == 12 else m + 1:02d}"
    return render_template("reports/analytics.html", rows=rows, k=k, peak=peak, top_lube=top_lube,
                           top_parties=top_parties, staff_n=staff_n, att_today=att_today, salary_paid=salary_paid,
                           exp_cat=exp_cat, stock_value=stock_value, low=low, outstanding=outstanding, fuels=fl,
                           month_name=f"{calendar.month_name[m]} {y}", ym=ym, prev_m=prev_m,
                           next_m=next_m if next_m <= today().strftime("%Y-%m") else None)


# ------------------------------------------------------------------ full business workbook with charts
@bp.route("/reports/business.xlsx")
@login_required("backup")
def business_xlsx():
    t = today()
    with db.tx() as cur:
        st = settings_map(cur)
        cur.execute("SELECT MIN(business_date) AS m FROM days")
        first = cur.fetchone()["m"] or t
        rows = range_summary(cur, max(first, t - timedelta(days=400)), t)
        parties = sorted(party_balances(cur), key=lambda r: -Decimal(r["balance"]))
        prods = _lube_stock(cur)
        from .staff import salary_sheet
        sheet = salary_sheet(cur, t.strftime("%Y-%m"))
        raw = []
        for tbl in BACKUP_TABLES:
            cur.execute(f"SELECT * FROM {tbl} ORDER BY 1")
            rr = cur.fetchall()
            cols = [c.name for c in cur.description if c.name != "password_hash"]
            raw.append((tbl, cols, [[(str(x[c]) if isinstance(x[c], (dict, list)) else x[c]) for c in cols] for x in rr]))
        audit(cur, "export", "backup", None, "Downloaded business workbook with charts")
    months = {}
    for r in rows:
        key = r["date"].strftime("%Y-%m")
        mm = months.setdefault(key, {"fuel": _Z, "lube": _Z, "credit": _Z, "received": _Z, "expenses": _Z})
        mm["fuel"] += r["fuel_total"]; mm["lube"] += r["lube_total"]; mm["credit"] += r["credit"]
        mm["received"] += r["payments"]; mm["expenses"] += r["expenses"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"
    ws["A1"] = f"{st.get('station_name', 'Sai Fuel Mart')} — business summary"; ws["A1"].font = Font(size=16, bold=True)
    ws["A2"] = f"Generated {t:%d %b %Y}"
    ws.append([])
    ws.append(["Total outstanding", float(sum((Decimal(p["balance"]) for p in parties if Decimal(p["balance"]) > 0), _Z))])
    ws.append(["Credit parties", len(parties)])
    ws.append(["Active staff", len(sheet["rows"])])
    ws.append(["Lube stock value", float(sum((Decimal(p["stock"]) * Decimal(p["sale_rate"]) for p in prods), _Z))])
    ws.append([])
    start_row = ws.max_row + 1
    ws.append(["Month", "Fuel sale", "Lube sale", "Credit given", "Payments received", "Expenses"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for key in sorted(months):
        mm = months[key]
        ws.append([key, float(mm["fuel"]), float(mm["lube"]), float(mm["credit"]), float(mm["received"]), float(mm["expenses"])])
    end_row = ws.max_row
    if end_row > start_row:
        ch = LineChart(); ch.title = "Monthly sale, credit and payments"; ch.y_axis.title = "₹"; ch.height = 9; ch.width = 22
        ch.add_data(Reference(ws, min_col=2, max_col=6, min_row=start_row, max_row=end_row), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=start_row + 1, max_row=end_row))
        ws.add_chart(ch, "H3")
    for col, w in zip("ABCDEF", (22, 16, 14, 16, 18, 14)):
        ws.column_dimensions[col].width = w
    wt = wb.create_sheet("Transport summary")
    write_sheet(wt, ["Party", "Opening", "Fuel credit", "Lube credit", "Received", "Discount", "Balance", "Last payment"],
                [[p["name"], p["opening_bal"], p["fuel_credit"], p["lube_credit"], p["received"], p["discounts"],
                  p["balance"], p["last_payment"]] for p in parties], (2, 3, 4, 5, 6, 7))
    top = [p for p in parties if Decimal(p["balance"]) > 0][:10]
    if top:
        pie = PieChart(); pie.title = "Outstanding by transporter (top 10)"; pie.height = 9; pie.width = 14
        pie.add_data(Reference(wt, min_col=7, min_row=1, max_row=len(top) + 1), titles_from_data=True)
        pie.set_categories(Reference(wt, min_col=1, min_row=2, max_row=len(top) + 1))
        wt.add_chart(pie, "J2")
    wl = wb.create_sheet("Lube summary")
    write_sheet(wl, ["Product", "Rate", "Stock", "Value", "Sold this month", "Reorder at"],
                [[p["name"], p["sale_rate"], p["stock"], Decimal(p["stock"]) * Decimal(p["sale_rate"]), p["sold_month"],
                  p["reorder_level"]] for p in prods], (2, 4))
    wsa = wb.create_sheet("Staff this month")
    write_sheet(wsa, ["Name", "Present", "Half", "Absent", "Leave", "Earned", "Advance", "Paid", "Still to pay"],
                [[r["staff"]["name"], r["att"].get("P", 0), r["att"].get("H", 0), r["att"].get("A", 0),
                  r["att"].get("L", 0), r["earned"], r["advance"], r["salary_paid"], r["balance"]] for r in sheet["rows"]],
                (6, 7, 8, 9))
    wd = wb.create_sheet("Daily sales")
    write_sheet(wd, ["Date", "Status", "Fuel", "Lube", "Total", "Digital", "Credit", "Received", "Expenses",
                     "Expected cash", "Counted", "Short/Excess"],
                [[r["date"], r["status"], r["fuel_total"], r["lube_total"], r["gross_sale"], r["digital"], r["credit"],
                  r["payments"], r["expenses"], r["expected_cash"], r["counted"], r["difference"]] for r in rows],
                tuple(range(3, 13)))
    if len(rows) > 1:
        bc = BarChart(); bc.title = "Daily sale"; bc.height = 8; bc.width = 24
        bc.add_data(Reference(wd, min_col=5, min_row=1, max_row=min(len(rows), 60) + 1), titles_from_data=True)
        bc.set_categories(Reference(wd, min_col=1, min_row=2, max_row=min(len(rows), 60) + 1))
        wd.add_chart(bc, "N2")
    for name, cols, data in raw:
        w = wb.create_sheet(f"data_{name}"[:31])
        write_sheet(w, cols, data, ())
    buf = BytesIO(); wb.save(buf); buf.seek(0)
    return _deliver_compat(buf, mimetype=XLSX, as_attachment=True, download_name=f"SaiFuelMart business {t:%Y-%m-%d}.xlsx")
